# -*- coding: utf-8 -*-
import sys
import os
sys.path.insert(0, os.path.abspath("."))

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
from src.screener.engine import run_screener

# Colour fills
GREEN  = PatternFill("solid", fgColor="C8E6C9")
RED    = PatternFill("solid", fgColor="FFCDD2")
YELLOW = PatternFill("solid", fgColor="FFF9C4")
HEADER = PatternFill("solid", fgColor="1A237E")
BENCH  = PatternFill("solid", fgColor="FFF8E1")

HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT   = Font(bold=True, size=13, color="1A237E")
NORMAL_FONT  = Font(size=10)

THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

# KPI columns to display
DISPLAY_COLS = [
    ("company_id",                "Ticker",        10),
    ("company_name",              "Company Name",  28),
    ("broad_sector",              "Sector",        20),
    ("composite_quality_score",   "Score (0-100)", 14),
    ("return_on_equity_pct",      "ROE %",         10),
    ("return_on_capital_pct",     "ROCE %",        10),
    ("net_profit_margin_pct",     "NPM %",         10),
    ("debt_to_equity",            "D/E",            8),
    ("interest_coverage",         "ICR",            8),
    ("free_cash_flow_cr",         "FCF (Cr)",      12),
    ("revenue_cagr_5yr",          "Rev CAGR 5yr",  14),
    ("pat_cagr_5yr",              "PAT CAGR 5yr",  14),
    ("operating_profit_margin_pct","OPM %",         10),
    ("pe_ratio",                  "P/E",            8),
    ("pb_ratio",                  "P/B",            8),
    ("dividend_yield_pct",        "Div Yield %",   12),
    ("capital_allocation_pattern","Cap Alloc",      18),
    ("icr_label",                 "ICR Label",      12),
]

# Thresholds for colour coding per preset
THRESHOLDS = {
    "quality_compounder": {
        "return_on_equity_pct":       ("min", 15),
        "debt_to_equity":             ("max", 1.0),
        "free_cash_flow_cr":          ("min", 0),
        "revenue_cagr_5yr":           ("min", 10),
    },
    "value_pick": {
        "pe_ratio":                   ("max", 35),
        "pb_ratio":                   ("max", 5.0),
        "debt_to_equity":             ("max", 2.0),
        "dividend_yield_pct":         ("min", 1),
    },
    "growth_accelerator": {
        "pat_cagr_5yr":               ("min", 20),
        "revenue_cagr_5yr":           ("min", 15),
        "debt_to_equity":             ("max", 2.0),
    },
    "dividend_champion": {
        "dividend_yield_pct":         ("min", 2),
        "dividend_payout_ratio_pct":  ("max", 80),
        "free_cash_flow_cr":          ("min", 0),
    },
    "debt_free_blue_chip": {
        "debt_to_equity":             ("max", 0.1),
        "return_on_equity_pct":       ("min", 12),
    },
    "turnaround_watch": {
        "revenue_cagr_3yr":           ("min", 10),
        "free_cash_flow_cr":          ("min", 0),
    },
}

PRESET_DESCRIPTIONS = {
    "quality_compounder":  "ROE>15%, D/E<1, FCF>0, Revenue CAGR 5yr>10%",
    "value_pick":          "P/E<35, P/B<5, D/E<2, Dividend Yield>1%",
    "growth_accelerator":  "PAT CAGR 5yr>20%, Revenue CAGR 5yr>15%, D/E<2",
    "dividend_champion":   "Dividend Yield>2%, Payout<80%, FCF>0",
    "debt_free_blue_chip": "D/E<0.1, ROE>12%, Sales>Rs.5000 Cr",
    "turnaround_watch":    "Revenue CAGR 3yr>10%, FCF positive",
}


def cell_colour(value, direction, threshold):
    """Return GREEN or RED fill based on threshold."""
    if value is None or (isinstance(value, float) and
                          value != value):
        return YELLOW
    try:
        val = float(value)
    except (ValueError, TypeError):
        return YELLOW
    if direction == "min":
        return GREEN if val >= threshold else RED
    else:
        return GREEN if val <= threshold else RED


def write_preset_sheet(wb, preset_name, results):
    """Write one sheet for a preset."""
    sheet_title = preset_name.replace("_", " ").title()
    ws = wb.create_sheet(title=sheet_title[:31])

    thresholds = THRESHOLDS.get(preset_name, {})

    # Row 1 — Preset title
    ws.merge_cells("A1:R1")
    ws["A1"] = sheet_title
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    # Row 2 — Description
    ws.merge_cells("A2:R2")
    ws["A2"] = PRESET_DESCRIPTIONS.get(preset_name, "")
    ws["A2"].font = Font(italic=True, size=10, color="455A64")
    ws["A2"].alignment = Alignment(horizontal="center")

    # Row 3 — Company count
    ws.merge_cells("A3:R3")
    ws["A3"] = f"Companies found: {len(results)}"
    ws["A3"].font = Font(bold=True, size=10)
    ws["A3"].alignment = Alignment(horizontal="center")

    # Row 4 — blank
    ws.row_dimensions[4].height = 8

    # Row 5 — Headers
    for col_idx, (col_key, col_label, col_width) in             enumerate(DISPLAY_COLS, start=1):
        cell = ws.cell(row=5, column=col_idx, value=col_label)
        cell.fill = HEADER
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center", wrap_text=True)
        cell.border = THIN
        ws.column_dimensions[
            get_column_letter(col_idx)].width = col_width

    ws.row_dimensions[5].height = 30

    # Data rows
    for row_idx, (_, data_row) in             enumerate(results.iterrows(), start=6):
        for col_idx, (col_key, col_label, _) in                 enumerate(DISPLAY_COLS, start=1):
            value = data_row.get(col_key, None)

            # Format numbers
            if isinstance(value, float):
                display_val = round(value, 2)
            else:
                display_val = value

            cell = ws.cell(
                row=row_idx, column=col_idx,
                value=display_val)
            cell.font = NORMAL_FONT
            cell.border = THIN
            cell.alignment = Alignment(horizontal="center")

            # Colour code threshold columns
            if col_key in thresholds:
                direction, threshold = thresholds[col_key]
                cell.fill = cell_colour(
                    value, direction, threshold)

        ws.row_dimensions[row_idx].height = 18

    # Freeze panes at row 6
    ws.freeze_panes = "A6"

    return ws


def generate_screener_excel():
    presets = [
        "quality_compounder",
        "value_pick",
        "growth_accelerator",
        "dividend_champion",
        "debt_free_blue_chip",
        "turnaround_watch",
    ]

    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    for preset in presets:
        print(f"Running {preset}...")
        results = run_screener(preset_name=preset)
        write_preset_sheet(wb, preset, results)
        print(f"  -> {len(results)} companies written")

    os.makedirs("output", exist_ok=True)
    output_path = "output/screener_output.xlsx"
    wb.save(output_path)
    print(f"\nSaved: {output_path}")
    return output_path


if __name__ == "__main__":
    path = generate_screener_excel()
    print("Done! Open output/screener_output.xlsx to view results.")
